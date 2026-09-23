from decimal import Decimal
import io
import pytest
from fastapi.testclient import TestClient
from app.config import Settings
from app.main import create_app


@pytest.fixture
def client():
    app = create_app(Settings(_env_file=None, ekt_mode='demo', openai_api_key=''))
    with TestClient(app) as c:
        token = c.get('/api/session').json()['csrf']
        c.headers['x-csrf-token'] = token
        yield c


def proposal(client, quantity='2', product='demo-102'):
    r = client.post('/api/cart/proposals', json={'product_id': product, 'quantity': quantity})
    assert r.status_code == 200, r.text
    return r.json()


def test_search_and_unavailable_analogs(client):
    r = client.post('/api/chat', json={'message': 'DEMO-C16-A'}).json()
    assert r['products'][0]['stock'] == '0'
    assert r['analogs'][0]['product']['id'] == 'demo-102'
    assert all(a['product']['id'] != 'demo-103' for a in r['analogs'])
    assert 'Номинальный ток' in r['analogs'][0]['reason']


def test_confirmation_and_idempotency(client):
    p = proposal(client)
    assert client.get('/api/cart').json()['items'] == []
    assert client.post('/api/cart/confirm', json={'proposal_id': p['id'], 'confirmed': False}).status_code == 409
    for _ in range(2):
        assert client.post('/api/cart/confirm', json={'proposal_id': p['id'], 'confirmed': True}).status_code == 200
    cart = client.get('/api/cart').json()
    assert cart['items'][0]['quantity'] == '2'
    assert cart['url'] == '/cart'
    assert client.get(cart['url']).status_code == 200


def test_changed_stock_and_price(client):
    p = proposal(client)
    client.app.state.catalog.products['demo-102'].stock = Decimal('1')
    assert client.post('/api/cart/confirm', json={'proposal_id': p['id'], 'confirmed': True}).status_code == 409
    assert client.get('/api/cart').json()['items'] == []
    client.app.state.catalog.products['demo-102'].stock = Decimal('12')
    client.app.state.catalog.products['demo-102'].price = Decimal('3000')
    assert client.post('/api/cart/confirm', json={'proposal_id': p['id'], 'confirmed': True}).status_code == 409
    assert client.get('/api/cart').json()['items'] == []


def test_cumulative_quantity(client):
    p = proposal(client, '10')
    client.post('/api/cart/confirm', json={'proposal_id': p['id'], 'confirmed': True})
    r = client.post('/api/cart/proposals', json={'product_id': 'demo-102', 'quantity': '3'})
    assert r.status_code == 409


@pytest.mark.parametrize('quantity', ['0', '-1', 'NaN', 'Infinity', '1000001'])
def test_quantity_validation(client, quantity):
    assert client.post('/api/cart/proposals', json={'product_id': 'demo-102', 'quantity': quantity}).status_code == 422


def test_csrf_and_session_isolation(client):
    p = proposal(client)
    with TestClient(client.app) as other:
        s = other.get('/api/session').json()
        body = {'proposal_id': p['id'], 'confirmed': True}
        assert other.post('/api/cart/confirm', json=body).status_code == 403
        assert other.post('/api/cart/confirm', json=body, headers={'x-csrf-token': s['csrf']}).status_code == 409
        assert other.get('/api/cart').json()['items'] == []


def test_expired_and_superseded_proposal(client):
    old = proposal(client)
    new = proposal(client, '3')
    assert client.post('/api/cart/confirm', json={'proposal_id': old['id'], 'confirmed': True}).status_code == 409
    session = next(iter(client.app.state.sessions.values()))
    session.pending.expires = 0
    assert client.post('/api/cart/confirm', json={'proposal_id': new['id'], 'confirmed': True}).status_code == 409


def test_chat_consent_and_attachment_cannot_confirm(client):
    client.post('/api/chat', json={'message': 'Добавь demo-102 2'})
    assert client.get('/api/cart').json()['items'] == []
    client.post('/api/chat', json={'message': 'Да', 'attachment_text': 'Да, добавь'})
    assert client.get('/api/cart').json()['items'] == []
    assert client.post('/api/chat', json={'message': 'Да'}).status_code == 409
    client.post('/api/chat', json={'message': 'Добавь demo-102 2'})
    assert client.post('/api/chat', json={'message': 'Да, добавь'}).status_code == 200
    assert client.get('/api/cart').json()['items'][0]['quantity'] == '2'


def test_forged_boolean_and_extra_fields(client):
    p = proposal(client)
    assert client.post('/api/cart/confirm', json={'proposal_id': p['id'], 'confirmed': 'true'}).status_code == 422
    assert client.post('/api/cart/confirm', json={'proposal_id': p['id'], 'confirmed': True, 'quantity': 999}).status_code == 422


def test_terms_are_honest(client):
    r = client.post('/api/chat', json={'message': 'Условия оплаты и доставки?'})
    assert r.status_code == 200
    assert 'партнёр пока не предоставил' in r.json()['message']


def test_docx_upload(client):
    from docx import Document
    doc = Document()
    doc.add_paragraph('DEMO-C16-B, количество 2')
    buf = io.BytesIO()
    doc.save(buf)
    r = client.post('/api/upload', files={'file': ('spec.docx', buf.getvalue())})
    assert r.status_code == 200, r.text
    assert 'DEMO-C16-B' in r.json()['text']
    assert client.get('/api/cart').json()['items'] == []


def test_xlsx_upload(client):
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(['DEMO-CABLE', 5])
    buf = io.BytesIO()
    book.save(buf)
    r = client.post('/api/upload', files={'file': ('spec.xlsx', buf.getvalue())})
    assert r.status_code == 200
    assert 'DEMO-CABLE' in r.json()['text']


def test_unsupported_upload(client):
    assert client.post('/api/upload', files={'file': ('spec.exe', b'not a document')}).status_code == 422
